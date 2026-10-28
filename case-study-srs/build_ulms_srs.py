"""Build reports for the ULMS Online software-requirements case study.

Usage from the repository root:
    python case-study-srs/build_ulms_srs.py
The script reads the requirements workbook and writes Markdown, CSV and JSON
artifacts under case-study-srs/output/.
"""
from __future__ import annotations

import csv
import json
import math
import sys
from collections import Counter
from pathlib import Path
from typing import Any
from types import SimpleNamespace

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "data" / "ULMS_Online_2026_Requirements_Demo.xlsx"
OUTPUT = Path(__file__).resolve().parent / "output"

sys.path.insert(0, str(ROOT))
from requirements_toolkit.requirement import (  # noqa: E402
    Priority, Requirement, RequirementSet, RequirementStatus, RequirementType,
)
from requirements_toolkit.srs_generator import SRSGenerator, SRSMetadata  # noqa: E402
from requirements_toolkit.traceability import (  # noqa: E402
    TraceabilityLink, TraceabilityMatrix, TraceabilityType,
)


def clean(value: Any) -> str:
    """Convert spreadsheet cell values into safe, readable strings."""
    if value is None:
        return ""
    if isinstance(value, float) and math.isnan(value):
        return ""
    return str(value).strip()


def read_sheet(workbook, sheet_name: str) -> list[dict[str, Any]]:
    """Read a worksheet as a list of dictionaries using its first row as headers."""
    if sheet_name not in workbook.sheetnames:
        return []
    sheet = workbook[sheet_name]
    rows = sheet.iter_rows(values_only=True)
    headers = [clean(v) for v in next(rows, ())]
    records = []
    for values in rows:
        if not any(v is not None and clean(v) for v in values):
            continue
        record = {
            headers[i]: (values[i] if i < len(values) else None)
            for i in range(len(headers)) if headers[i]
        }
        records.append({key: clean(value) for key, value in record.items()})
    return records


def enum_by_value(enum_type, value, default):
    text = clean(value).casefold()
    for member in enum_type:
        if member.value.casefold() == text or member.name.casefold() == text.replace("-", "_").replace(" ", "_"):
            return member
    return default


def build_requirement_set(records: list[dict[str, Any]]) -> RequirementSet:
    """Convert workbook requirement rows into toolkit Requirement objects."""
    result = RequirementSet(name="University Library Management System Online")
    for row in records:
        req_id = clean(row.get("id"))
        title = clean(row.get("title"))
        statement = clean(row.get("description"))
        if not req_id or not title or not statement:
            continue
        raw_deps = clean(row.get("dependencies"))
        raw_tags = clean(row.get("tags"))
        deps = [item.strip() for item in raw_deps.replace(",", ";").split(";") if item.strip()]
        tags = [item.strip() for item in raw_tags.replace(",", ";").split(";") if item.strip()]
        requirement = Requirement(
            req_id=req_id,
            title=title,
            statement=statement,
            req_type=enum_by_value(RequirementType, row.get("type"), RequirementType.FUNCTIONAL),
            priority=enum_by_value(Priority, row.get("priority"), Priority.SHOULD),
            status=enum_by_value(RequirementStatus, row.get("status"), RequirementStatus.PROPOSED),
            source=clean(row.get("source")),
            rationale=clean(row.get("rationale")),
            acceptance_criteria=clean(row.get("acceptance_criteria")),
            verification_method=clean(row.get("verification_method")) or "Test",
            complexity=clean(row.get("complexity")),
            dependencies=deps,
            tags=tags,
        )
        result.add(requirement)
    return result


def main() -> dict[str, Any]:
    if not WORKBOOK.exists():
        raise FileNotFoundError(
            f"Requirements workbook not found: {WORKBOOK}\n"
            "Place ULMS_Online_2026_Requirements_Demo.xlsx in the data/ folder."
        )
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    try:
        sheet_names = [
            "project", "stakeholders", "requirements", "nfrs", "actors",
            "use_cases", "change_requests", "hazards", "elicitation_sessions",
            "process_activities", "test_cases", "test_case_analysis",
        ]
        tables = {name: read_sheet(workbook, name) for name in sheet_names}
    finally:
        workbook.close()

    requirements = build_requirement_set(tables["requirements"])
    OUTPUT.mkdir(parents=True, exist_ok=True)

    metadata = SRSMetadata(
        project_name="University Library Management System Online (ULMS Online)",
        document_title="Software Requirements Specification",
        version="1.0",
        status="Draft - for review",
        organization="University Library",
    )
    nfr_objects = [
        SimpleNamespace(
            nfr_id=row.get("id", ""),
            title=row.get("title", ""),
            category=row.get("category", ""),
            statement=row.get("description", ""),
            priority=row.get("priority", ""),
            metric=row.get("metric", ""),
            target_value=row.get("target_value", ""),
            measurement_unit="",
        )
        for row in tables["nfrs"]
    ]
    stakeholder_objects = [
        SimpleNamespace(
            stakeholder_id=row.get("id", ""),
            name=row.get("name", ""),
            role=row.get("role", ""),
            stakeholder_type=row.get("organization", ""),
            responsibilities=row.get("notes", ""),
            needs=row.get("notes", ""),
        )
        for row in tables["stakeholders"]
    ]
    use_case_objects = [
        SimpleNamespace(
            use_case_id=row.get("id", ""),
            name=row.get("title", ""),
            description=row.get("goal", ""),
            primary_actor=row.get("primary_actor", ""),
            priority="",
            status="",
            preconditions=[row.get("preconditions", "")] if row.get("preconditions") else [],
            main_flow=[step.strip() for step in row.get("main_flow", "").split(";") if step.strip()],
            alternative_flows=[step.strip() for step in row.get("alternative_flows", "").split(";") if step.strip()],
            postconditions=[row.get("postconditions", "")] if row.get("postconditions") else [],
        )
        for row in tables["use_cases"]
    ]

    matrix = TraceabilityMatrix(name="ULMS Workbook Traceability")
    for row in tables["use_cases"]:
        use_case_id = clean(row.get("id"))
        for req_id in [x.strip() for x in clean(row.get("related_requirements")).replace(",", ";").split(";") if x.strip()]:
            if requirements.get(req_id):
                matrix.add_link(TraceabilityLink(
                    req_id, use_case_id,
                    TraceabilityType.REQUIREMENT_TO_USE_CASE,
                    f"Related to use case: {clean(row.get('title'))}",
                ))
    for row in tables["test_cases"]:
        test_id = clean(row.get("id"))
        for req_id in [x.strip() for x in clean(row.get("related_requirements")).replace(",", ";").split(";") if x.strip()]:
            if requirements.get(req_id):
                matrix.add_link(TraceabilityLink(
                    req_id, test_id,
                    TraceabilityType.REQUIREMENT_TO_TEST_CASE,
                    f"Verified by test case: {clean(row.get('title'))}",
                ))

    generator = SRSGenerator(
        metadata=metadata,
        requirements=requirements,
        nfrs=nfr_objects,
        stakeholders=stakeholder_objects,
        use_cases=use_case_objects,
        traceability=matrix,
    )
    srs_path = generator.save(str(OUTPUT / "ULMS_SRS.md"))

    # Export the source requirement records as CSV without losing workbook columns.
    req_rows = tables["requirements"]
    csv_path = OUTPUT / "ULMS_requirements.csv"
    fieldnames = list(dict.fromkeys(key for row in req_rows for key in row.keys()))
    with csv_path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames or ["id", "title", "description"])
        writer.writeheader()
        writer.writerows(req_rows)

    counts = {name: len(rows) for name, rows in tables.items()}
    dashboard = {
        "project": "University Library Management System Online",
        "workbook": WORKBOOK.name,
        "generated_at_note": "Generated locally when this script was run.",
        "sheet_record_counts": counts,
        "requirement_type_counts": dict(Counter(
            clean(row.get("type")) or "Unspecified" for row in req_rows
        )),
        "requirement_priority_counts": dict(Counter(
            clean(row.get("priority")) or "Unspecified" for row in req_rows
        )),
        "output_files": [str(srs_path.relative_to(ROOT)), str(csv_path.relative_to(ROOT))],
    }
    traceability_path = OUTPUT / "ULMS_traceability.json"
    traceability_path.write_text(
        json.dumps([link.to_dict() for link in matrix.links], ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    dashboard["traceability_links"] = len(matrix.links)
    dashboard["output_files"].append(str(traceability_path.relative_to(ROOT)))
    dashboard_path = OUTPUT / "ULMS_dashboard.json"
    dashboard_path.write_text(json.dumps(dashboard, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"Workbook: {WORKBOOK}")
    print(f"Requirements loaded into toolkit: {len(requirements.requirements)}")
    print("Worksheet record counts:")
    for name, count in counts.items():
        print(f"  {name}: {count}")
    print(f"SRS: {srs_path}")
    print(f"CSV: {csv_path}")
    print(f"Dashboard: {dashboard_path}")
    print(f"Traceability links: {len(matrix.links)} ({traceability_path})")
    return dashboard


if __name__ == "__main__":
    main()
