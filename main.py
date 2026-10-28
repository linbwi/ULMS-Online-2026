"""Command-line entry point for the ULMS Online Requirements Engineering toolkit.

Run from the project root with `python main.py`.
This is separate from the Flask website (`python app.py`).
"""
from __future__ import annotations

import importlib.util
import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook
from requirements_toolkit.validation import RequirementValidator

ROOT = Path(__file__).resolve().parent
BUILDER_PATH = ROOT / "case-study-srs" / "build_ulms_srs.py"
WORKBOOK = ROOT / "data" / "ULMS_Online_2026_Requirements_Demo.xlsx"
OUTPUT = ROOT / "case-study-srs" / "output"


def load_builder():
    spec = importlib.util.spec_from_file_location("ulms_srs_builder", BUILDER_PATH)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Cannot load SRS builder: {BUILDER_PATH}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def main() -> int:
    print("=" * 72)
    print("ULMS ONLINE — REQUIREMENTS ENGINEERING TOOLKIT")
    print("=" * 72)

    if not WORKBOOK.exists():
        print(f"ERROR: Requirements workbook not found: {WORKBOOK}")
        return 1

    builder = load_builder()
    print("\nGenerating reports from the current workbook...")
    dashboard = builder.main()
    req_set = builder.build_requirement_set(
        # The builder's reader normalizes spreadsheet values consistently.
        _read_requirements(builder, WORKBOOK)
    )
    reqs = list(req_set)

    print("\n=== Requirements ===")
    print(f"Total requirements loaded into toolkit: {len(reqs)}")
    for req in reqs:
        print(f"{req.req_id} | {req.priority.value} | {req.title}")

    print("\n=== Testability / Quality Check ===")
    report = RequirementValidator().validate_requirement_set(req_set)
    if not report.issues:
        print("No potential issues detected by automated checks.")
    else:
        for issue in report.issues:
            print(str(issue))
    summary = report.summary()
    print(
        f"Summary: {summary['total_issues']} issue(s), "
        f"{summary['errors']} error(s), {summary['warnings']} warning(s), "
        f"{summary['infos']} info message(s)"
    )
    print("Note: these heuristic checks support, but do not replace, human review.")

    print("\n=== Priority Breakdown ===")
    priority_counts = Counter(req.priority.value for req in reqs)
    for priority in ("Must", "Should", "Could", "Won't"):
        print(f"{priority}: {priority_counts.get(priority, 0)}")

    print("\n=== Traceability / Coverage Report ===")
    trace_path = OUTPUT / "ULMS_traceability.json"
    links = json.loads(trace_path.read_text(encoding="utf-8")) if trace_path.exists() else []
    req_ids = {req.req_id for req in reqs}
    test_linked_ids = set()
    use_case_linked_ids = set()
    for link in links:
        link_type = str(link.get("link_type", link.get("type", ""))).upper()
        source = str(link.get("source_id", link.get("source", "")))
        target = str(link.get("target_id", link.get("target", "")))
        candidates = {source, target}
        for req_id in req_ids.intersection(candidates):
            other = (candidates - {req_id})
            if any(token in link_type for token in ("TEST", "VERIFICATION")) or any(
                token in " ".join(other).upper() for token in ("TC-", "TEST-")
            ):
                test_linked_ids.add(req_id)
            if "USE_CASE" in link_type or any(token in " ".join(other).upper() for token in ("UC-", "USE-CASE")):
                use_case_linked_ids.add(req_id)

    # Fallback to dashboard/source records if serialized link schema differs.
    if not test_linked_ids:
        test_linked_ids = _count_related_ids(WORKBOOK, "test_cases", "related_requirements", req_ids)
    if not use_case_linked_ids:
        use_case_linked_ids = _count_related_ids(WORKBOOK, "use_cases", "related_requirements", req_ids)

    test_coverage = len(test_linked_ids) / len(req_ids) * 100 if req_ids else 0
    use_case_coverage = len(use_case_linked_ids) / len(req_ids) * 100 if req_ids else 0
    print(f"Traceability links: {len(links)}")
    print(f"Requirements linked to use cases: {len(use_case_linked_ids)} / {len(req_ids)} ({use_case_coverage:.1f}%)")
    print(f"Requirements linked to test cases: {len(test_linked_ids)} / {len(req_ids)} ({test_coverage:.1f}%)")
    print(f"Requirements without test-case links: {len(req_ids - test_linked_ids)}")
    print("A test-case link does not mean the test has passed.")

    print("\n=== Workbook Sheet Summary ===")
    for name, count in dashboard.get("sheet_record_counts", {}).items():
        print(f"{name}: {count} record(s)")

    print("\n=== Generated Artifacts ===")
    for name in ("ULMS_SRS.md", "ULMS_requirements.csv", "ULMS_traceability.json", "ULMS_dashboard.json"):
        path = OUTPUT / name
        print(f"{name}: {'OK' if path.exists() else 'NOT GENERATED'}")
    print("\nFinished. To launch the website, run `python app.py` separately.")
    return 0


def _read_requirements(builder, workbook_path):
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        return builder.read_sheet(workbook, "requirements")
    finally:
        workbook.close()


def _count_related_ids(workbook_path, sheet_name, column_name, requirement_ids):
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        rows = load_builder().read_sheet(workbook, sheet_name)
    finally:
        workbook.close()
    linked = set()
    for row in rows:
        raw = str(row.get(column_name, "") or "")
        for item in raw.replace(",", ";").split(";"):
            item = item.strip()
            if item in requirement_ids:
                linked.add(item)
    return linked


if __name__ == "__main__":
    raise SystemExit(main())
